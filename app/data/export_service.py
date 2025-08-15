# -*- coding: utf-8 -*-
"""Utilities to export database tables to CSV or Excel files."""
from __future__ import annotations

import csv
from typing import List, Tuple

from . import db


ALLOWED_TABLES = (
    "clientes",
    "dispositivos",
    "inventario",
    "repuestos",
    "reparaciones",
    "usuarios",
)


def _fetch_table(table: str) -> Tuple[List[str], List[Tuple]]:
    """Return column names and all rows for the given table."""
    if table not in ALLOWED_TABLES:
        raise ValueError(f"Invalid table name: {table}")
    conn = db._ensure_conn()
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM {table}")
    rows = cur.fetchall()
    headers = [col[0] for col in cur.description]
    return headers, rows


def export_table_to_csv(table: str, filepath: str) -> None:
    """Export the given table to a CSV file."""
    headers, rows = _fetch_table(table)
    with open(filepath, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


def export_table_to_excel(table: str, filepath: str) -> None:
    """Export the given table to an Excel file using openpyxl."""
    headers, rows = _fetch_table(table)
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - defensive
        raise RuntimeError("openpyxl is required for Excel export") from exc
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(filepath)


def import_table_from_csv(table: str, filepath: str) -> None:
    """Import rows from a CSV file into the given table.

    Existing rows in the table will be removed before import.
    """
    if table not in ALLOWED_TABLES:
        raise ValueError(f"Invalid table name: {table}")
    with open(filepath, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        headers = next(reader)
        rows = [[None if val == "" else val for val in row] for row in reader]
    conn = db._ensure_conn()
    cur = conn.cursor()
    cur.execute(f"DELETE FROM {table}")
    placeholders = ", ".join("?" for _ in headers)
    cur.executemany(
        f"INSERT INTO {table} ({', '.join(headers)}) VALUES ({placeholders})",
        rows,
    )
    conn.commit()
